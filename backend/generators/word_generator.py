from docx import Document
from docx.shared import Pt


def create_word(title, content, output_path):

    document = Document()

    title_paragraph = document.add_heading(
        title,
        level=1
    )

    for item in content:

        if isinstance(item, dict):

            heading = item.get("heading")
            text = item.get("text")

            if heading:
                document.add_heading(
                    heading,
                    level=2
                )

            if text:
                document.add_paragraph(text)

        else:
            document.add_paragraph(str(item))

    document.save(output_path)

    return output_path
