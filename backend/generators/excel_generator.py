from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def create_excel(title, columns, rows, output_path):

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "PageGenieCraft"

    sheet["A1"] = title
    sheet["A1"].font = Font(
        size=16,
        bold=True
    )

    for column_index, column in enumerate(columns, 1):

        cell = sheet.cell(
            row=3,
            column=column_index
        )

        cell.value = column
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            start_color="4F46E5",
            end_color="4F46E5",
            fill_type="solid"
        )

    for row_index, row in enumerate(rows, 4):

        for column_index, value in enumerate(row, 1):

            sheet.cell(
                row=row_index,
                column=column_index
            ).value = value

    for column in sheet.columns:

        max_length = 0

        column_letter = column[0].column_letter

        for cell in column:

            if cell.value:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        sheet.column_dimensions[
            column_letter
        ].width = min(max_length + 2, 50)

    workbook.save(output_path)

    return output_path
