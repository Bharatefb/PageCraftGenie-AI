from pptx import Presentation


def create_powerpoint(title, slides, output_path):

    presentation = Presentation()

    # Title slide
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[0]
    )

    slide.shapes.title.text = title

    for item in slides:

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[1]
        )

        slide.shapes.title.text = item.get(
            "title",
            "Slide"
        )

        body = slide.placeholders[1]

        content = item.get(
            "content",
            []
        )

        body.text = "\n".join(
            str(x) for x in content
        )

    presentation.save(output_path)

    return output_path
